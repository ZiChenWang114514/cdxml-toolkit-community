"""Read local structure tables without evaluating formulas or inferring measurements."""
from pathlib import Path
import csv
import hashlib
import json
import re
from rdkit import Chem
from ..chemistry_semantics import require_supported, semantic_key


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def local_path(base, value):
    path = Path(value).expanduser()
    return path.resolve() if path.is_absolute() else (Path(base)/path).resolve()


def _excel_text(value, number_format, location):
    if value is None:
        return ''
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return str(value)
    if number_format in {'General', '@'}:
        return str(value)
    match = re.fullmatch(r'0(?:\.(0+))?(%?)(?:E\+00)?(\s*"[^"]*")?', number_format)
    if not match:
        raise ValueError(f'{location}: unsupported numeric format {number_format!r}; supply an explicit text value to preserve its display')
    decimals = len(match[1] or '')
    number = value * 100 if match[2] else value
    rendered = format(number, f'.{decimals}' + ('E' if 'E+00' in number_format else 'f'))
    return rendered + match[2] + (match[3] or '').replace('"', '')


def table_rows(path, sheet=None):
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.reader(stream)
            values = list(reader)
    elif path.suffix.lower() == '.xlsx':
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError('XLSX requires the publication optional dependency (openpyxl).') from exc
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet is None and len(formulas.sheetnames) != 1:
                raise ValueError('Select sheet explicitly: ' + ', '.join(formulas.sheetnames))
            chosen = sheet or formulas.sheetnames[0]
            if chosen not in formulas.sheetnames:
                raise ValueError(f'Unknown worksheet: {chosen}')
            values = []
            for row_index, (frow, crow) in enumerate(zip(formulas[chosen].iter_rows(), cached[chosen].iter_rows()), 1):
                row = []
                for column_index, (cell, result) in enumerate(zip(frow, crow), 1):
                    location = f'{chosen}!{openpyxl.utils.get_column_letter(column_index)}{row_index}'
                    if cell.data_type == 'f' and result.value is None:
                        raise ValueError(f'{location}: formula has no cached value; calculate and save in Excel')
                    row.append(_excel_text(result.value, cell.number_format,
                                           location))
                values.append(row)
        finally:
            formulas.close(); cached.close()
    else:
        raise ValueError('Table must be CSV or XLSX')
    if not values:
        raise ValueError('Table has no header row')
    header = [str(v).strip() for v in values[0]]
    if not all(header) or len(set(header)) != len(header):
        raise ValueError('Header row contains empty or duplicate column names')
    rows = []
    for number, cells in enumerate(values[1:], 2):
        if not any(str(c).strip() for c in cells):
            continue
        if len(cells) != len(header):
            raise ValueError(f'row {number}: expected {len(header)} columns, got {len(cells)}')
        rows.append((number, dict(zip(header, cells))))
    return header, rows


def load_records(spec, base):
    table = local_path(base, spec['table'])
    header, rows = table_rows(table, spec.get('sheet'))
    if not 1 <= len(rows) <= 1000:
        raise ValueError('Table must contain 1..1000 compound rows')
    fields = spec.get('fields', [])
    if len({f['column'] for f in fields}) != len(fields):
        raise ValueError('Duplicate display field')
    id_column = spec.get('id_column', 'compound_id')
    structure = spec.get('structure', {'format':'smiles', 'column':'smiles'})
    fmt = structure.get('format', 'smiles')
    if fmt not in {'smiles','mol','sdf'}:
        raise ValueError('structure.format must be smiles, mol, or sdf')
    required = {id_column, *(f['column'] for f in fields)}
    if spec.get('group_column'): required.add(spec['group_column'])
    if spec.get('label_column'): required.add(spec['label_column'])
    if fmt != 'sdf': required.add(structure.get('column', 'smiles' if fmt=='smiles' else 'structure_file'))
    if required - set(header):
        raise ValueError('Missing columns: ' + ', '.join(sorted(required-set(header))))
    sources = {str(table): sha256(table)}
    sdf = {}
    if fmt == 'sdf':
        path = local_path(base, structure['file']); sources[str(path)] = sha256(path)
        prop = structure.get('id_property', id_column)
        for index, mol in enumerate(Chem.SDMolSupplier(str(path), removeHs=False), 1):
            if mol is None or not mol.HasProp(prop):
                raise ValueError(f'SDF record {index}: invalid structure or missing {prop}')
            sdf.setdefault(mol.GetProp(prop).strip(), []).append(mol)
    records = []
    seen = set()
    for row_number, row in rows:
        cid = row[id_column].strip()
        if not cid or cid in seen:
            raise ValueError(f'row {row_number}: empty or duplicate compound_id {cid!r}')
        seen.add(cid)
        try:
            if fmt == 'smiles':
                mol = Chem.MolFromSmiles(row[structure.get('column','smiles')])
            elif fmt == 'mol':
                path = local_path(table.parent, row[structure.get('column','structure_file')])
                sources[str(path)] = sha256(path)
                mol = Chem.MolFromMolFile(str(path), removeHs=False)
            else:
                candidates = sdf.get(cid, [])
                if len(candidates) != 1:
                    raise ValueError(f'{len(candidates)} SDF candidates for {cid}')
                mol = candidates[0]
            if mol is None or not mol.GetNumAtoms():
                raise ValueError('invalid or empty structure')
            if mol.GetNumAtoms() > 1000:
                raise ValueError('structure exceeds 1000 atoms')
            require_supported(mol)
        except Exception as exc:
            raise ValueError(f'row {row_number} ({cid}): {exc}') from exc
        records.append({'compound_id':cid, 'row':row_number,
                        'label':row.get(spec.get('label_column','label'), '') or cid,
                        'group':row.get(spec.get('group_column',''), ''),
                        # New layout must not import old CX coordinates/wedge directives.
                        'smiles':semantic_key(mol), 'identity':semantic_key(mol),
                        'values':{f['column']:row[f['column']] for f in fields}})
    return records, sources


def field_text(record, field):
    value = record['values'][field['column']]
    # A missing value is visible, never replaced with a numeric zero.
    rendered = value if value != '' else '—'
    if value != '' and field.get('unit'): rendered += ' ' + field['unit']
    label = field.get('label', field['column'])
    return f'{label}: {rendered}' if label else rendered


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
